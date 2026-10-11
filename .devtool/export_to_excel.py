"""
Script para exportar los resultados individuales y grupales de MSEED a bitacora.xlsx.
"""
import os
import glob
import json
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_PATH = r"d:\2026\ARDASYSMIC\ARDASYSMIC_26\Trabajos\Promart-CityMall\bitacora.xlsx"
INDIV_RESULTS_DIR = r"d:\2026\ARDASYSMIC\ARDASYSMIC_26\Trabajos\Promart-CityMall\.devtool\results_individual"
GROUP_RESULTS_DIR = r"d:\2026\ARDASYSMIC\ARDASYSMIC_26\Trabajos\Promart-CityMall\.devtool\results_group"

def get_sesame_class(f0, a0):
    if f0 <= 0 or a0 <= 0:
        return "Desconocido"
    elif a0 < 1.0:
        return "E (H/V menor que 1)"
    elif a0 < 1.5:
        return "D (H/V de amplitud baja)"
    elif a0 > 2.0:
        if f0 > 1.0:
            return "A (pico destacado, f0 > 1 Hz)"
        else:
            return "B (pico destacado, f0 ≤ 1 Hz)"
    else:
        return "C (pico poco destacado)"

def extract_data(results_dir, is_group=False):
    rows = []
    
    # Expresión regular para extraer el nombre del sitio de la carpeta
    site_regex = re.compile(r"^(N\d+(?:\s+actualizado)?)(?:_|$)")
    
    subdirs = sorted([d for d in glob.glob(os.path.join(results_dir, "N*")) if os.path.isdir(d)])
    
    for sd in subdirs:
        dir_name = os.path.basename(sd)
        
        match = site_regex.match(dir_name)
        sitio_name = match.group(1) if match else dir_name
        
        json_matches = glob.glob(os.path.join(sd, "*__resultados.json"))
        if not json_matches:
            continue
            
        with open(json_matches[0], "r", encoding="utf-8") as f:
            data = json.load(f)
            
        if is_group:
            # Para grupos, source puede ser una carpeta
            mseed_file = "GRUPO_" + sitio_name
        else:
            mseed_file = os.path.basename(data.get("source", ""))
            
        traces = data.get("traces", [])
        tr0 = traces[0] if traces else {}
        stats = data.get("statistics", {})
        
        f0 = data.get("f0_hz", 0.0)
        a0 = data.get("a0", 0.0)
        t0_s = data.get("t0_s", 0.0)
        kg = data.get("kg", 0.0)
        reliable = data.get("reliable", False)
        sesame = get_sesame_class(f0, a0)
        
        row = {
            "sitio": sitio_name,
            "mseed": mseed_file,
            "estacion": tr0.get("station", ""),
            "red": tr0.get("network", ""),
            "canales": len(traces),
            "fs_hz": tr0.get("sample_rate_hz", 0.0),
            "muestras": tr0.get("sample_count", 0),
            "duracion_s": tr0.get("duration_s", 0.0),
            "tiempo_inicio": tr0.get("start_time", ""),
            "ventanas": data.get("num_windows", 0),
            "ventana_s": data.get("window_length_s", 0.0),
            "f0_hz": f0,
            "t0_s": t0_s,
            "a0": a0,
            "kg": kg,
            "sesame": sesame,
            "confiable": reliable,
            "num_cycles": data.get("num_cycles", 0.0),
            "f0_low_hz": stats.get("f0_low_hz"),
            "f0_high_hz": stats.get("f0_high_hz"),
            "f0_median_hz": stats.get("f0_median_hz"),
            "t_mean_hz": stats.get("t_mean_hz"),
            "periodo_low_s": stats.get("period_low_s"),
            "periodo_median_s": stats.get("period_median_s"),
            "periodo_high_s": stats.get("period_high_s"),
            "unidades": data.get("units_status", ""),
            "stationxml": os.path.basename(data.get("stationxml", "") or ""),
            "stationxml_completo": data.get("stationxml_complete", False),
            "latitud": tr0.get("latitude"),
            "longitud": tr0.get("longitude"),
            "elevacion_m": tr0.get("elevation_m"),
        }
        rows.append(row)
    return rows

def format_sheet(ws, rows):
    headers = [
        ("sitio", "Sitio"),
        ("mseed", "Archivo / Grupo"),
        ("estacion", "Estación"),
        ("canales", "Canales"),
        ("fs_hz", "Fs (Hz)"),
        ("duracion_s", "Duración (s)"),
        ("ventanas", "Ventanas"),
        ("ventana_s", "Ventana (s)"),
        ("f0_hz", "f0 (Hz)"),
        ("t0_s", "T0 (s)"),
        ("a0", "A0"),
        ("kg", "Kg"),
        ("sesame", "Clasificación SESAME"),
        ("confiable", "Confiable SESAME"),
        ("num_cycles", "Ciclos Nc"),
        ("f0_low_hz", "f0 Min (Hz)"),
        ("f0_high_hz", "f0 Max (Hz)"),
        ("f0_median_hz", "f0 Mediana (Hz)"),
        ("t_mean_hz", "T Medio (Hz)"),
        ("periodo_low_s", "T Min (s)"),
        ("periodo_median_s", "T Mediana (s)"),
        ("periodo_high_s", "T Max (s)"),
        ("tiempo_inicio", "Hora Inicio"),
        ("unidades", "Unidades"),
        ("stationxml", "StationXML"),
        ("stationxml_completo", "StationXML Completo"),
        ("latitud", "Latitud"),
        ("longitud", "Longitud"),
        ("elevacion_m", "Elevación (m)"),
    ]
    
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    highlight_font = Font(name="Calibri", size=10, bold=True, color="002060")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    right_align = Alignment(horizontal='right', vertical='center')
    
    for col_idx, (key, title) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    for row_idx, rdata in enumerate(rows, 2):
        for col_idx, (key, _) in enumerate(headers, 1):
            val = rdata.get(key)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            
            if key in ["sitio", "estacion", "canales", "ventanas", "confiable", "stationxml_completo"]:
                cell.alignment = center_align
            elif key in ["mseed", "tiempo_inicio", "unidades", "stationxml", "sesame"]:
                cell.alignment = left_align
            else:
                cell.alignment = right_align
                
            if key in ["fs_hz", "duracion_s", "ventana_s", "num_cycles", "elevacion_m"]:
                cell.number_format = "#,##0.0"
            elif key in ["f0_hz", "f0_low_hz", "f0_high_hz", "f0_median_hz", "t_mean_hz"]:
                cell.number_format = "0.0000"
            elif key in ["t0_s", "periodo_low_s", "periodo_median_s", "periodo_high_s"]:
                cell.number_format = "0.0000"
                cell.font = highlight_font
            elif key in ["a0", "kg"]:
                cell.number_format = "#,##0.000"
            elif key in ["latitud", "longitud"]:
                cell.number_format = "0.0000000"
                
    for col_idx, (key, title) in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(title)
        for rdata in rows:
            v_str = str(rdata.get(key, ""))
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

def main():
    print("Extrayendo resultados individuales...")
    indiv_rows = extract_data(INDIV_RESULTS_DIR, is_group=False)
    
    print("Extrayendo resultados grupales...")
    group_rows = extract_data(GROUP_RESULTS_DIR, is_group=True)
    
    wb = openpyxl.load_workbook(EXCEL_PATH)
    
    # Manejar Hoja Individual
    indiv_sheet_name = "Analisis Individual"
    if indiv_sheet_name in wb.sheetnames:
        wb.remove(wb[indiv_sheet_name])
    elif "Hoja2" in wb.sheetnames:
        wb.remove(wb["Hoja2"]) # Eliminar hoja vieja
        
    ws_indiv = wb.create_sheet(title=indiv_sheet_name)
    format_sheet(ws_indiv, indiv_rows)
    
    # Manejar Hoja Grupal
    group_sheet_name = "Analisis Grupal"
    if group_sheet_name in wb.sheetnames:
        wb.remove(wb[group_sheet_name])
        
    ws_group = wb.create_sheet(title=group_sheet_name)
    format_sheet(ws_group, group_rows)
        
    wb.save(EXCEL_PATH)
    print(f"Guardado exitosamente en: {EXCEL_PATH}")
    print(f"  - {len(indiv_rows)} registros individuales en '{indiv_sheet_name}'")
    print(f"  - {len(group_rows)} registros grupales en '{group_sheet_name}'")

if __name__ == "__main__":
    main()
